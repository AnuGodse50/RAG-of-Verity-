"""Loaders for JSON, CSV, and Excel data without schema assumptions."""
from pathlib import Path
import csv, json, hashlib
from core.models import Document
from core.interfaces import Loader
from loaders.registry import register_loader
def _sid(p: Path): return hashlib.sha256(str(p.resolve()).encode()).hexdigest()
@register_loader(".csv", ".tsv")
class CSVLoader(Loader):
    def load(self, source: str) -> list[Document]:
        p=Path(source); delim="\t" if p.suffix.lower()==".tsv" else ","
        with p.open(encoding="utf-8-sig", newline="") as f: rows=list(csv.DictReader(f, delimiter=delim))
        sid=_sid(p); return [Document(id=f"{sid}:row{i}", text=json.dumps(row, ensure_ascii=False), metadata={"source":str(p),"source_id":sid,"type":"tabular","row":i}) for i,row in enumerate(rows,1)]
@register_loader(".xlsx", ".xlsm", ".xls")
class ExcelLoader(Loader):
    def load(self, source: str) -> list[Document]:
        p=Path(source);sid=_sid(p)
        if p.suffix.lower()==".xls":
            import pandas as pd
            out=[]
            for sheet,frame in pd.read_excel(p,sheet_name=None).items():
                for i,row in enumerate(frame.where(frame.notna(),None).to_dict(orient="records"),1):
                    out.append(Document(id=f"{sid}:{sheet}:{i}",text=json.dumps(row,default=str,ensure_ascii=False),metadata={"source":str(p),"source_id":sid,"type":"tabular","sheet":sheet,"row":i}))
            return out
        from openpyxl import load_workbook
        wb=load_workbook(p, read_only=True, data_only=True); out=[]
        for ws in wb.worksheets:
            rows=ws.iter_rows(values_only=True); headers=[str(x or f"column_{i+1}") for i,x in enumerate(next(rows,()))]
            for i,row in enumerate(rows,1):
                values={headers[j]:v for j,v in enumerate(row) if j<len(headers) and v is not None}
                out.append(Document(id=f"{sid}:{ws.title}:{i}",text=json.dumps(values,default=str,ensure_ascii=False),metadata={"source":str(p),"source_id":sid,"type":"tabular","sheet":ws.title,"row":i}))
        return out
@register_loader(".json", ".jsonl")
class JSONLoader(Loader):
    def load(self, source: str) -> list[Document]:
        p=Path(source); sid=_sid(p); raw=p.read_text(encoding="utf-8")
        values=[json.loads(line) for line in raw.splitlines() if line.strip()] if p.suffix==".jsonl" else [json.loads(raw)]
        return [Document(id=f"{sid}:item{i}",text=json.dumps(value,ensure_ascii=False,indent=2),metadata={"source":str(p),"source_id":sid,"type":"json","item":i}) for i,value in enumerate(values)]
