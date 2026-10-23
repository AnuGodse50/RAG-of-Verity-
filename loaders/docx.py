"""Microsoft Word document loader."""
from pathlib import Path
import hashlib
from docx import Document as WordDocument
from core.models import Document
from core.interfaces import Loader
from loaders.registry import register_loader
@register_loader(".docx")
class DOCXLoader(Loader):
    def load(self,source:str)->list[Document]:
        path=Path(source);doc=WordDocument(path);parts=[]
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        sid=hashlib.sha256(str(path.resolve()).encode()).hexdigest()
        return [Document(id=sid,text="\n".join(parts),metadata={"source":str(path),"source_id":sid,"type":"docx"})]
